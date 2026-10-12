import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain.positions import ROSTER_MAX, Position
from app.models import Franchise
from app.roster_import.writers import COLUMNS


@pytest.fixture
def team_id(franchise: Franchise) -> int:
    return franchise.team_id


def _export(client: TestClient, team_id: int, fmt: str):
    response = client.get(f"/api/teams/{team_id}/roster-export.{fmt}")
    assert response.status_code == 200
    return response


def _snapshot(client: TestClient, team_id: int) -> list[tuple]:
    """Everything an export should preserve, independent of database ids."""
    rows = []
    for summary in client.get(f"/api/teams/{team_id}/roster").json():
        p = client.get(f"/api/players/{summary['id']}").json()
        rows.append((
            p["first_name"], p["last_name"], p["position"], p["jersey_number"], p["age"],
            p["height_in"], p["weight_lb"], p["college"], p["years_pro"], p["overall"],
            p["potential"], p["development_trait"], p["injury_status"],
            p["contract"]["salary"], p["contract"]["end_year"],
            tuple(sorted(p["ratings"].items())),
        ))
    return sorted(rows)


class TestRosterExport:
    def test_csv_layout_and_filename(self, client: TestClient, team_id: int):
        response = _export(client, team_id, "csv")
        assert response.headers["content-type"].startswith("text/csv")
        assert 'filename="HFD-roster-2026.csv"' in response.headers["content-disposition"]
        lines = response.text.splitlines()
        assert lines[0].split(",") == COLUMNS
        assert len(lines) == 1 + ROSTER_MAX
        # Grouped by position in depth-chart order: quarterbacks first, punter last.
        assert lines[1].split(",")[2] == "QB"
        assert lines[-1].split(",")[2] == "P"

    def test_xlsx_has_frozen_bold_header(self, client: TestClient, team_id: int):
        content = _export(client, team_id, "xlsx").content
        sheet = load_workbook(io.BytesIO(content)).worksheets[0]
        assert sheet.title == "HFD 2026"
        assert [c.value for c in sheet[1]] == COLUMNS
        assert sheet["A1"].font.bold
        assert sheet.freeze_panes == "A2"
        assert sheet.max_row == 1 + ROSTER_MAX

    def test_json_includes_team_metadata(self, client: TestClient, team_id: int):
        body = json.loads(_export(client, team_id, "json").content)
        assert body["team"] == "Hartford Ironclads"
        assert body["season"] == 2026
        assert len(body["players"]) == ROSTER_MAX
        assert set(body["players"][0]) <= set(COLUMNS)

    @pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
    def test_export_reimports_unchanged(self, client: TestClient, team_id: int, fmt: str):
        roster = client.get(f"/api/teams/{team_id}/roster").json()
        # Awkward values that must survive the round trip.
        client.patch(f"/api/players/{roster[0]['id']}",
                     json={"last_name": "O'Neil, Jr.", "injury_status": "injured_reserve",
                           "development_trait": "elite"})
        client.patch(f"/api/players/{roster[1]['id']}", json={"college": 'The "U" of Somewhere'})
        before = _snapshot(client, team_id)

        content = _export(client, team_id, fmt).content
        preview = client.post(f"/api/teams/{team_id}/roster-import/preview",
                              files={"file": (f"export.{fmt}", content)},
                              data={"mode": "replace"}).json()
        assert preview["summary"] == {"total": ROSTER_MAX, "valid": ROSTER_MAX,
                                      "warnings": 0, "errors": 0}
        result = client.post(f"/api/teams/{team_id}/roster-import/commit",
                             files={"file": (f"export.{fmt}", content)},
                             data={"mode": "replace"}).json()
        assert result["imported"] == ROSTER_MAX

        assert _snapshot(client, team_id) == before

    def test_short_contract_exports_at_least_one_year(self, client: TestClient, team_id: int,
                                                      db):
        from app.models import Player
        player = db.query(Player).filter(Player.team_id == team_id,
                                         Player.position == Position.K.value).one()
        player.active_contract.start_year = 2020
        player.active_contract.years = 1  # expired in 2020
        db.flush()
        lines = _export(client, team_id, "csv").text.splitlines()
        kicker = next(line.split(",") for line in lines if line.split(",")[2] == "K")
        assert kicker[COLUMNS.index("contract_years")] == "1"

    def test_unknown_team_and_format(self, client: TestClient, team_id: int):
        assert client.get("/api/teams/999999/roster-export.csv").status_code == 404
        assert client.get(f"/api/teams/{team_id}/roster-export.pdf").status_code == 422


def test_browser_can_read_export_filename(client: TestClient, franchise: Franchise):
    # The browser downloads through the web app's own /api (same origin), so it can read
    # the filename without any CORS exposure.
    response = client.get(f"/api/teams/{franchise.team_id}/roster-export.csv")
    assert response.headers["content-disposition"].startswith("attachment; filename=")

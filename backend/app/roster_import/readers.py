"""Read an uploaded roster file into raw rows: a list of {column: value} dicts.

No validation happens here beyond the file being readable; see rows.py for that.
"""

import csv
import io
import json
import zipfile
from enum import StrEnum
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500

RawRow = dict[str, Any]


class FileFormat(StrEnum):
    CSV = "csv"
    XLSX = "xlsx"
    JSON = "json"


class ImportFileError(Exception):
    """The file as a whole can't be read (wrong type, corrupt, too big, empty)."""


_EXTENSIONS = {".csv": FileFormat.CSV, ".tsv": FileFormat.CSV, ".txt": FileFormat.CSV,
               ".xlsx": FileFormat.XLSX, ".json": FileFormat.JSON}


def detect_format(filename: str) -> FileFormat:
    name = filename.lower()
    if name.endswith(".xls"):
        raise ImportFileError("Legacy .xls files aren't supported; save as .xlsx or CSV.")
    for extension, fmt in _EXTENSIONS.items():
        if name.endswith(extension):
            return fmt
    raise ImportFileError("Unsupported file type. Upload a .csv, .xlsx, or .json file.")


def read_rows(content: bytes, fmt: FileFormat) -> list[RawRow]:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"File is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if not content.strip():
        raise ImportFileError("File is empty.")
    readers = {FileFormat.CSV: _read_csv, FileFormat.XLSX: _read_xlsx,
               FileFormat.JSON: _read_json}
    rows = [row for row in readers[fmt](content) if any(_present(v) for v in row.values())]
    if not rows:
        raise ImportFileError("No player rows found.")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Too many rows ({len(rows)}); the limit is {MAX_ROWS}.")
    return rows


def _present(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _read_csv(content: bytes) -> list[RawRow]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    sample = text[:4096]
    # Only the delimiter is sniffed. The sniffer's quoting guesses are unreliable (it can
    # misread doubled quotes inside a field), so standard CSV quoting is always used.
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        raise ImportFileError("CSV has no header row.")
    # DictReader puts overflow cells under None; drop them.
    return [{k: v for k, v in row.items() if k is not None} for row in reader]


def _read_xlsx(content: bytes) -> list[RawRow]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        raise ImportFileError("Couldn't open the Excel file. Is it a valid .xlsx?") from exc
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header or not any(_present(h) for h in header):
        raise ImportFileError("The first sheet has no header row.")
    names = [str(h).strip() if h is not None else "" for h in header]
    result = [{name: value for name, value in zip(names, row, strict=False) if name}
              for row in rows]
    workbook.close()
    return result


def _read_json(content: bytes) -> list[RawRow]:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ImportFileError(f"Invalid JSON: {exc}") from exc
    if isinstance(data, dict):
        data = data.get("players")
    if not isinstance(data, list):
        raise ImportFileError('JSON must be a list of players or {"players": [...]}.')
    rows: list[RawRow] = []
    for item in data:
        if not isinstance(item, dict):
            raise ImportFileError("Every JSON player must be an object.")
        row = {k: v for k, v in item.items() if not isinstance(v, dict)}
        # Accept the API's own shape: a nested "ratings" object.
        ratings = item.get("ratings")
        if isinstance(ratings, dict):
            row.update(ratings)
        contract = item.get("contract")
        if isinstance(contract, dict):
            row.setdefault("salary", contract.get("salary"))
            row.setdefault("contract_years", contract.get("years"))
        rows.append(row)
    return rows
